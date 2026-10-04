import openpyxl
path = r"data\NGA_West2_Flatfile_RotD50_d005_public_version.xlsx"
wb = openpyxl.load_workbook(path, read_only=True)
print("Sheets:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print("\n==", name, "| rows:", ws.max_row, "| columns:", ws.max_column)
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= 4:
            break
        print(i, [c for c in row][:12])
